import os
import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.services.firebase_db import firestore_db

def create_excel():
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    
    # -------------------------------------------------------------
    # 1. Colors & Styles
    # -------------------------------------------------------------
    FONT_FAMILY = "Segoe UI"
    
    title_font = Font(name=FONT_FAMILY, size=16, bold=True, color="1F497D")
    subtitle_font = Font(name=FONT_FAMILY, size=10, italic=True, color="595959")
    section_font = Font(name=FONT_FAMILY, size=12, bold=True, color="1F497D")
    header_font = Font(name=FONT_FAMILY, size=11, bold=True, color="FFFFFF")
    bold_font = Font(name=FONT_FAMILY, size=10, bold=True, color="000000")
    data_font = Font(name=FONT_FAMILY, size=10, color="000000")
    kpi_val_font = Font(name=FONT_FAMILY, size=18, bold=True, color="1F497D")
    kpi_lbl_font = Font(name=FONT_FAMILY, size=9, bold=True, color="595959")
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    teacher_header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    student_header_fill = PatternFill(start_color="333f48", end_color="333f48", fill_type="solid")
    class_header_fill = PatternFill(start_color="548235", end_color="548235", fill_type="solid")
    summary_card_fill = PatternFill(start_color="F2F4F7", end_color="F2F4F7", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    admin_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    
    thin_border_side = Side(border_style="thin", color="D9D9D9")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # -------------------------------------------------------------
    # Fetch Data
    # -------------------------------------------------------------
    students = []
    if firestore_db.is_connected:
        try:
            students = firestore_db.get_students_list()
        except Exception as e:
            print(f"Error fetching from firestore: {e}")
            
    with open('data/app_state.json') as f:
        local_state = json.load(f)
    
    existing_sids = {s.get("student_id") for s in students}
    for s in local_state.get("students", []):
        if s.get("student_id") and s.get("student_id") not in existing_sids:
            students.append(s)
            existing_sids.add(s.get("student_id"))

    # Teachers list
    teachers = [
        {
            "teacher_id": "teacher_bio101",
            "name": "Dr. Evelyn Reed",
            "department": "Biology & Life Sciences",
            "class_code": "bio101",
            "pin": "teacher123",
            "role": "Teacher / Class Supervisor",
            "class_name": "Bio 101: General Biology & Ecology",
            "status": "Active"
        },
        {
            "teacher_id": "teacher_env201",
            "name": "Prof. Marcus Green",
            "department": "Environmental Sciences",
            "class_code": "env201",
            "pin": "teacher123",
            "role": "Teacher / Class Supervisor",
            "class_name": "Env 201: Climate, Carbon & Soil Systems",
            "status": "Active"
        },
        {
            "teacher_id": "teacher_cs301",
            "name": "Ms. Clara Patel",
            "department": "Computer Science & IT",
            "class_code": "cs301",
            "pin": "teacher123",
            "role": "Teacher / Class Supervisor",
            "class_name": "CS 301: Cloud Computing & Systems",
            "status": "Active"
        },
        {
            "teacher_id": "teacher_chem102",
            "name": "Dr. Arthur Thorne",
            "department": "Chemistry",
            "class_code": "chem102",
            "pin": "teacher123",
            "role": "Teacher / Class Supervisor",
            "class_name": "Chem 102: Organic Chemistry & Sustainability",
            "status": "Active"
        },
        {
            "teacher_id": "teacher_earth204",
            "name": "Mr. Samuel Woods",
            "department": "Earth & Planetary Sciences",
            "class_code": "earth204",
            "pin": "teacher123",
            "role": "Teacher / Class Supervisor",
            "class_name": "Earth 204: Global Biosphere & Forestry",
            "status": "Active"
        }
    ]

    admins = [
        {
            "admin_id": "admin",
            "name": "System Administrator",
            "department": "IT Operations & School Administration",
            "class_code": "All Classes",
            "pin": "admin2026",
            "secondary_pin": "admin123",
            "role": "Global Administrator",
            "status": "Active"
        }
    ]

    classes = [
        {"code": "bio101", "name": "Bio 101: General Biology & Ecology", "instructor": "Dr. Evelyn Reed", "dept": "Life Sciences"},
        {"code": "env201", "name": "Env 201: Climate, Carbon & Soil Systems", "instructor": "Prof. Marcus Green", "dept": "Environmental Sciences"},
        {"code": "cs301", "name": "CS 301: Cloud Computing & Systems", "instructor": "Ms. Clara Patel", "dept": "Computer Science"},
        {"code": "chem102", "name": "Chem 102: Organic Chemistry & Sustainability", "instructor": "Dr. Arthur Thorne", "dept": "Chemistry"},
        {"code": "earth204", "name": "Earth 204: Global Biosphere & Forestry", "instructor": "Mr. Samuel Woods", "dept": "Earth Sciences"},
        {"code": "class-a", "name": "Class A: Pilot Demonstration Group", "instructor": "Staff Demo Account", "dept": "Interdisciplinary"}
    ]

    teacher_by_class = {t["class_code"]: t["name"] for t in teachers}

    # =============================================================
    # SHEET 1: Summary Overview
    # =============================================================
    ws_sum = wb.create_sheet(title="Overview & Credentials")
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum["B2"] = "DATA-TO-DIRT: Master User Directory & Access Guide"
    ws_sum["B2"].font = title_font
    ws_sum["B3"] = "Generated from Production Cloud Database • Includes Student, Teacher, and Administrator Accounts"
    ws_sum["B3"].font = subtitle_font
    
    # KPI Summary Cards (row 5 label, row 6 value)
    kpis = [
        ("B", "C", "TOTAL REGISTERED USERS", f"=COUNTA('All Students'!B4:B{len(students)+3}) + COUNTA('Teachers & Staff'!B4:B{len(teachers)+len(admins)+3})"),
        ("D", "E", "STUDENT ACCOUNTS", f"=COUNTA('All Students'!B4:B{len(students)+3})"),
        ("F", "G", "FACULTY & ADMINS", f"=COUNTA('Teachers & Staff'!B4:B{len(teachers)+len(admins)+3})"),
        ("H", "I", "ACTIVE CLASSROOMS", f"=COUNTA('Classes & Departments'!A4:A{len(classes)+3})")
    ]
    
    for c1, c2, lbl, formula in kpis:
        # Merge top row for label
        ws_sum.merge_cells(f"{c1}5:{c2}5")
        cell_lbl = ws_sum[f"{c1}5"]
        cell_lbl.value = lbl
        cell_lbl.font = kpi_lbl_font
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
        cell_lbl.fill = summary_card_fill
        
        # Merge bottom row for value
        ws_sum.merge_cells(f"{c1}6:{c2}6")
        cell_val = ws_sum[f"{c1}6"]
        cell_val.value = formula
        cell_val.font = kpi_val_font
        cell_val.alignment = Alignment(horizontal="center", vertical="center")
        cell_val.fill = summary_card_fill
        
        # Borders
        for r in (5, 6):
            for c_idx in [ord(c1)-ord('A')+1, ord(c2)-ord('A')+1]:
                ws_sum.cell(row=r, column=c_idx).border = cell_border
                ws_sum.cell(row=r, column=c_idx).fill = summary_card_fill

    ws_sum.row_dimensions[5].height = 20
    ws_sum.row_dimensions[6].height = 28

    # Quick Access Instructions Table
    ws_sum["B9"] = "Role-Based Authentication Protocols"
    ws_sum["B9"].font = section_font
    
    inst_headers = ["User Role", "Portal Login Method", "Identifier (Username)", "Security PIN / Password", "Authorized Portal Views"]
    for col_idx, h in enumerate(inst_headers, start=2):
        cell = ws_sum.cell(row=11, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = cell_border
        
    inst_rows = [
        ("Student (Learner)", "Select 'Student' tab on login page", "Student ID (e.g. STU-001 or STU-BIO-002)", "Class Code (e.g. bio101, env201)", "Student Dashboard (AI Scanner, Tree Evolution, Classroom Leaderboard)"),
        ("Teacher (Staff)", "Select 'Staff' tab or click teacher quick-pills", "Teacher ID (e.g. teacher_bio101)", "teacher123", "Teacher Queue Review, Student Directory Creation, Classroom Analytics"),
        ("Administrator", "Select 'Admin' tab or click Admin auto-fill", "admin", "admin2026 (or admin123)", "Eco-Fund & Tree Donations, Teacher Approvals, Global Telemetry & Security Logs")
    ]
    
    for row_idx, r_data in enumerate(inst_rows, start=12):
        for col_idx, val in enumerate(r_data, start=2):
            cell = ws_sum.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            if row_idx % 2 == 1:
                cell.fill = zebra_fill
    
    ws_sum.column_dimensions["A"].width = 4
    ws_sum.column_dimensions["B"].width = 20
    ws_sum.column_dimensions["C"].width = 30
    ws_sum.column_dimensions["D"].width = 28
    ws_sum.column_dimensions["E"].width = 25
    ws_sum.column_dimensions["F"].width = 35
    ws_sum.column_dimensions["G"].width = 15
    ws_sum.column_dimensions["H"].width = 15
    ws_sum.column_dimensions["I"].width = 15
    ws_sum.row_dimensions[11].height = 24
    
    # =============================================================
    # SHEET 2: Teachers & Staff Directory
    # =============================================================
    ws_tea = wb.create_sheet(title="Teachers & Staff")
    ws_tea.views.sheetView[0].showGridLines = True
    ws_tea.freeze_panes = "A4"
    
    ws_tea["A1"] = "Faculty & Administrator Accounts"
    ws_tea["A1"].font = title_font
    ws_tea["A2"] = "Complete credentials and class management assignments"
    ws_tea["A2"].font = subtitle_font
    
    tea_headers = [
        "Role Type", "Account ID", "Full Name / Title", "Department", 
        "Assigned Class Code", "Classroom Name", "Login Password / PIN", "Access Rights & Responsibilities", "Status"
    ]
    
    for col_idx, h in enumerate(tea_headers, start=1):
        cell = ws_tea.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = teacher_header_fill
        cell.alignment = Alignment(horizontal="center" if "Code" in h or "Status" in h or "PIN" in h else "left", vertical="center", wrap_text=True)
        cell.border = cell_border
        
    curr_row = 4
    # Admin first
    for a in admins:
        row_vals = [
            a["role"], a["admin_id"], a["name"], a["department"], 
            a["class_code"], "All Institutional Classrooms", a["pin"], 
            "Full System Access: Tree-Nation donations, fund management, and global audit logs", a["status"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_tea.cell(row=curr_row, column=col_idx, value=val)
            cell.font = bold_font if col_idx <= 3 else data_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 5, 7, 9) else "left", vertical="center")
            cell.fill = admin_fill
        curr_row += 1

    # Teachers
    for t in teachers:
        row_vals = [
            t["role"], t["teacher_id"], t["name"], t["department"], 
            t["class_code"], t["class_name"], t["pin"], 
            f"Review deletion requests for {t['class_code']}, generate new student IDs, monitor eco-points", t["status"]
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_tea.cell(row=curr_row, column=col_idx, value=val)
            cell.font = bold_font if col_idx in (2, 3) else data_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 2, 5, 7, 9) else "left", vertical="center")
            if curr_row % 2 == 1:
                cell.fill = zebra_fill
        curr_row += 1

    ws_tea.row_dimensions[3].height = 24
    for col in ws_tea.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_tea.column_dimensions[col_letter].width = max(max_len + 3, 14)

    # =============================================================
    # SHEET 3: Students / Test Users Directory
    # =============================================================
    ws_stu = wb.create_sheet(title="All Students")
    ws_stu.views.sheetView[0].showGridLines = True
    ws_stu.freeze_panes = "A4"
    
    ws_stu["A1"] = "Student & Learner Accounts Directory"
    ws_stu["A1"].font = title_font
    ws_stu["A2"] = f"Total Registered: {len(students)} Students • Synced with Live Cloud Firestore & Local Datastore"
    ws_stu["A2"].font = subtitle_font
    
    stu_headers = [
        "Student ID (Username)", "Full Name", "Class Code (PIN)", "Assigned Course", 
        "Faculty Advisor", "Eco-Points", "Data Cleared (Bytes)", "Trees Funded", "Account Created At", "Created By"
    ]
    
    for col_idx, h in enumerate(stu_headers, start=1):
        cell = ws_stu.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = student_header_fill
        cell.alignment = Alignment(horizontal="center" if "Code" in h or "ID" in h or "Points" in h or "Trees" in h else "left", vertical="center", wrap_text=True)
        cell.border = cell_border

    # Sort students by class code, then student_id
    sorted_students = sorted(students, key=lambda x: (x.get("class_code", "").lower(), x.get("student_id", "")))

    for idx, s in enumerate(sorted_students, start=4):
        cc = s.get("class_code", "bio101").lower()
        adv = teacher_by_class.get(cc, s.get("created_by") or "Dr. Evelyn Reed")
        course_name = next((c["name"] for c in classes if c["code"] == cc), cc.upper())
        
        row_vals = [
            s.get("student_id") or f"STU-{idx:03d}",
            s.get("name") or "Student",
            cc,
            course_name,
            adv,
            s.get("eco_points", 0),
            s.get("bytes_cleared", 0),
            s.get("trees", 0),
            s.get("created_at") or "2026-10-22",
            s.get("created_by") or "Teacher"
        ]
        
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_stu.cell(row=idx, column=col_idx, value=val)
            cell.font = bold_font if col_idx in (1, 2) else data_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 3, 6, 7, 8, 9) else "left", vertical="center")
            if idx % 2 == 1:
                cell.fill = zebra_fill
                
            # Number formats
            if col_idx in (6, 8):
                cell.number_format = "#,##0"
            elif col_idx == 7:
                cell.number_format = "#,##0"

    ws_stu.row_dimensions[3].height = 24
    for col in ws_stu.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_stu.column_dimensions[col_letter].width = max(min(max_len + 3, 38), 12)

    # =============================================================
    # SHEET 4: Classes & Curriculum
    # =============================================================
    ws_cls = wb.create_sheet(title="Classes & Departments")
    ws_cls.views.sheetView[0].showGridLines = True
    ws_cls.freeze_panes = "A4"
    
    ws_cls["A1"] = "Classroom & Curriculum Master List"
    ws_cls["A1"].font = title_font
    ws_cls["A2"] = "Classroom codes used for student authentication and teacher scoping"
    ws_cls["A2"].font = subtitle_font
    
    cls_headers = [
        "Class Code", "Course Title", "Faculty in Charge", "Academic Department", "Student Enrollment", "Student Login Instruction"
    ]
    
    for col_idx, h in enumerate(cls_headers, start=1):
        cell = ws_cls.cell(row=3, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = class_header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 5) else "left", vertical="center", wrap_text=True)
        cell.border = cell_border

    for idx, c in enumerate(classes, start=4):
        enrollment_formula = f'=COUNTIF(\'All Students\'!C4:C{len(sorted_students)+3}, "{c["code"]}")'
        row_vals = [
            c["code"],
            c["name"],
            c["instructor"],
            c["dept"],
            enrollment_formula,
            f'Students in this class enter Student ID as username and "{c["code"]}" as PIN'
        ]
        for col_idx, val in enumerate(row_vals, start=1):
            cell = ws_cls.cell(row=idx, column=col_idx, value=val)
            cell.font = bold_font if col_idx <= 2 else data_font
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center" if col_idx in (1, 5) else "left", vertical="center")
            if idx % 2 == 1:
                cell.fill = zebra_fill
            if col_idx == 5:
                cell.number_format = "#,##0"

    ws_cls.row_dimensions[3].height = 24
    for col in ws_cls.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_cls.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Remove initial default sheet if still present
    if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_sheet)

    # Save to locations
    out_paths = [
        "/Users/narayan/.gemini/antigravity/scratch/data-to-dirt/Data_to_Dirt_Test_Users_and_Teachers.xlsx",
        "/Users/narayan/.gemini/antigravity/scratch/data-to-dirt/src/ui/Data_to_Dirt_Test_Users_and_Teachers.xlsx",
        "/Users/narayan/.gemini/antigravity/brain/5281e7b6-b8b6-4185-adfd-aade0da8e189/Data_to_Dirt_Test_Users_and_Teachers.xlsx"
    ]
    
    for p in out_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        wb.save(p)
        print(f"[✓] Saved workbook to: {p}")

if __name__ == "__main__":
    create_excel()
